from datetime import datetime, timedelta
from pathlib import Path
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, func

from monitoring import api
from monitoring.store import engine, initialize, observations, now, jobs, save_readings, set_source, timestamp
from monitoring.collect import import_excel, record, sync_welang_api
from monitoring.queries import history, snapshot
from monitoring.exports import run_export

@pytest.mark.parametrize("station,sensor,reference", [
    ("lokasi_1_hang_tuah", "distance", 500),
    ("lokasi_2_kalibokor", "distance1", 500),
    ("lokasi_3_rumah_pompa_pucang", "distance2", 530),
    ("lokasi_4_pln_menur", "distance1", 300),
])
def test_backend_calibration_preserves_raw_archive(tmp_path, monkeypatch, station, sensor, reference):
    from api.main import app
    db = engine('sqlite:///' + str(tmp_path / 'archive.db'))
    initialize(db)
    # Old archive rows may still classify the distance column as water_level.
    readings = [record(station, sensor, 'water_level', f'2026-10-18T12:0{index}:00+07:00', value,
        'cm', 'mysql:test', str(index), 'unverified') for index, value in enumerate((100, None, reference + 1, 0))]
    save_readings(db, readings)
    monkeypatch.setattr(api, '_db', db)
    with TestClient(app) as client:
        response = client.get(f'/v1/stations/{station}/history?hours=24').json()
        series = response['series'][0]
        assert series['referenceCm'] == reference
        assert series['measurement'] == 'distance_to_surface'
        assert [point['value'] for point in series['points']] == [100, None, reference + 1, 0]
        assert [point['waterLevelCm'] for point in series['points']] == [reference - 100, None, 0, reference]
        current = client.get(f'/v1/stations/{station}/snapshot').json()
        assert current['referenceCm'] == reference
        legacy = client.get('/surabaya/snapshot').json()
        active = next(item for item in legacy['stations'] if item['id'] == station)
        signal = next(item for item in active['sensors'] if item['id'] == sensor)
        assert signal['referenceCm'] == reference
        assert signal['valueCm'] == 0 and signal['waterLevelCm'] == reference
    with db.connect() as conn:
        stored = conn.execute(select(observations.c.value).order_by(observations.c.observed_at)).scalars().all()
    assert stored == [100, None, reference + 1, 0]

def test_excel_idempotency_and_mapping(tmp_path):
    db = engine('sqlite:///' + str(tmp_path / 'archive.db'))
    initialize(db)
    file = tmp_path / 'welang.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['DateTime', 'RC', 'RL', 'LP', 'LD'])
    ws.append([datetime(2023, 9, 15, 12), 1.5, 2.5, 0.107, 0.315])
    ws.append(['bad-time', 0, 0, 1, 2])
    wb.save(file)
    result = import_excel(db, file)
    assert result == {'rows': 1, 'skipped': 1}
    assert import_excel(db, file)['duplicate'] is True
    with db.connect() as conn:
        rows = conn.execute(select(observations)).mappings().all()
    assert len(rows) == 4
    purwodadi = next(r for r in rows if r['station'] == 'welang-purwodadi')
    assert purwodadi['unit'] == 'm' and purwodadi['value'] == 0.107
    assert purwodadi['observed_at'] == datetime(2023, 9, 15, 5)
    assert snapshot(db, 'welang-dhompo')['state'] == 'historical'

def test_api_priority_and_outage_preserve_excel(tmp_path, monkeypatch):
    db = engine('sqlite:///' + str(tmp_path / 'archive.db'))
    initialize(db)
    when = '2023-09-15T12:00:00+07:00'
    excel = record('welang-dhompo', 'water_level', 'water_level', when, .3, 'm', 'excel:welang', 'a', 'confirmed_by_owner')
    upstream = record('welang-dhompo', 'water_level', 'water_level', when, .5, 'm', 'api:welang', '1', 'provider_declared')
    save_readings(db, [excel, upstream])
    save_readings(db, [upstream])
    rows = history(db, 'welang-dhompo', timestamp(when) - timedelta(hours=1), timestamp(when))
    assert rows['series'][0]['points'][0]['value'] == .5
    monkeypatch.setenv('WELANG_API_URL', 'http://provider.invalid')
    def fail(*args, **kwargs):
        raise TimeoutError()
    monkeypatch.setattr('monitoring.collect.urlopen', fail)
    sync_welang_api(db)
    assert snapshot(db, 'welang-dhompo')['series'][0]['points'][-1]['value'] == .5

def test_export_months_raw_records_and_empty_sheets(tmp_path, monkeypatch):
    from monitoring import exports
    monkeypatch.setattr(exports, 'EXPORT_DIR', tmp_path)
    db = engine('sqlite:///' + str(tmp_path / 'archive.db'))
    initialize(db)
    save_readings(db, [record('welang-dhompo', 'water_level', 'water_level', '2023-09-15T12:00:00+07:00', .315, 'm', 'excel:welang', '1', 'confirmed_by_owner')])
    with db.begin() as conn:
        conn.execute(jobs.insert().values(id='test', station='welang-dhompo', status='processing', created_at=now(), expires_at=now() + timedelta(hours=24)))
    run_export(db, {'id': 'test', 'station': 'welang-dhompo'})
    wb = openpyxl.load_workbook(tmp_path / 'test.xlsx', read_only=True, data_only=True)
    assert len(wb.sheetnames) == 12
    assert wb.sheetnames[0] == '2022-10' and wb.sheetnames[-1] == '2023-09'
    assert wb['2023-09']['D6'].value == .315
    assert wb['2023-09']['A6'].value == datetime(2023, 9, 15, 12)
    assert wb['2022-10']['A6'].value == 'Tidak ada data pada bulan ini.'
    wb.close()

def test_routes_missing_station_no_data_and_export_lifecycle(tmp_path, monkeypatch):
    from api.main import app
    db = engine('sqlite:///' + str(tmp_path / 'archive.db'))
    initialize(db)
    monkeypatch.setattr(api, '_db', db)
    with TestClient(app) as client:
        assert client.get('/health/ready').status_code == 200
        assert client.get('/v1/stations/unknown/snapshot').status_code == 404
        assert client.get('/v1/stations/welang-dhompo/history').json()['series'] == []
        assert client.post('/v1/stations/welang-dhompo/exports').status_code == 409
        assert client.post('/predict').status_code == 503
        assert client.get('/v1/exports/not-a-job/file').status_code == 404
        assert client.get('/v1/capabilities').json()['laravel']['enabled'] is True

def test_welang_cm_display_keeps_original_export_unit(tmp_path):
    db = engine('sqlite:///' + str(tmp_path / 'archive.db'))
    initialize(db)
    save_readings(db, [record('welang-dhompo', 'water_level', 'water_level', '2023-09-15T12:00:00+07:00', 31.5, 'cm', 'api:welang', '1', 'provider_declared')])
    assert snapshot(db, 'welang-dhompo')['series'][0]['unit'] == 'm'
    assert snapshot(db, 'welang-dhompo')['series'][0]['points'][-1]['value'] == .315
    with db.connect() as conn:
        original = conn.execute(select(observations)).mappings().one()
    assert original['value'] == 31.5 and original['unit'] == 'cm'

def test_surabaya_independent_live_and_backfill_cursors(tmp_path, monkeypatch):
    import monitoring.collect as collector
    from monitoring.store import read_source
    db = engine('sqlite:///' + str(tmp_path / 'archive.db'))
    initialize(db)
    config = tmp_path / 'sources.yaml'
    config.write_text('host: example\nusername: readonly\npassword: test\n')
    monkeypatch.setenv('SURABAYA_DATABASE_CONFIG', str(config))
    monkeypatch.setattr(collector, 'live_config', lambda: {'locations': [
        {'key': 'test', 'database': 'dbtest', 'water_level': {'table': 'esp1', 'id_column': 'id', 'timestamp_column': 'waktu', 'columns': ['distance1']}}
    ]})
    class Connection:
        def __init__(self): self.calls = []; self.last = ''; self.cycle = 0
        def cursor(self): return self
        def __enter__(self): return self
        def __exit__(self, *args): pass
        def execute(self, sql, args=None): self.last = sql; self.calls.append((sql, args))
        def fetchall(self):
            if 'AND' in self.last: return [{'id': 1, 'waktu': datetime(2026, 1, 1, 12), 'distance1': 100.0}]
            if 'DESC' in self.last: return [{'id': 10, 'waktu': datetime(2026, 10, 17, 12), 'distance1': 123.0}]
            if 'WHERE' in self.last: return [{'id': 11, 'waktu': datetime(2026, 10, 17, 12, 1), 'distance1': 124.0}]
            return []
        def rollback(self): pass
        def close(self): pass
    connection = Connection()
    monkeypatch.setattr(collector.pymysql, 'connect', lambda **kwargs: connection)
    collector.sync_surabaya(db)
    state = read_source(db, 'mysql:dbtest.esp1')
    assert state['cursor'] == 10
    assert json.loads(state['detail'])['backfill_complete'] is True
    collector.sync_surabaya(db)
    assert read_source(db, 'mysql:dbtest.esp1')['cursor'] == 11
    with db.connect() as conn:
        assert conn.execute(select(func.count()).select_from(observations)).scalar() == 3
    assert any('TRANSACTION READ ONLY' in sql for sql, _ in connection.calls)
    assert not any(sql.lstrip().upper().startswith(('INSERT', 'UPDATE', 'DELETE')) for sql, _ in connection.calls)
